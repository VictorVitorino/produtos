"""Relatório do ciclo: HTML autocontido (pt-BR) e exportação XLSX, montados a partir do registro."""
import json
import re
import zipfile
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .comparacao import COMPONENTES, FAIXAS, NOMES
from .util import D, agora, brl, num, relogio_fixo

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
PILL = {"acima da materialidade": "r", "acima da faixa": "a", "dentro da faixa": "g", "classe própria": "b",
        "recebida": "g", "parcial": "a", "faltante": "r", "aceito": "g", "rejeitado": "r", "pendente": "a",
        "permitida": "g", "devolvida": "a", "fechado": "g", "aberto": "a"}
PARAM_NOMES = {"tarifa_movimentacao": "Tarifa de movimentação (R$/mil m³)", "tarifa_capacidade": "Tarifa de capacidade",
               "preco_gas_uso_sistema": "Preço do gás de uso do sistema (R$/mil m³)"}
CSS = """
:root{--bg:#f5f6f8;--card:#fff;--ink:#18212d;--muted:#5d6878;--line:#e2e6ec;--accent:#1d5f8a;
--r:#b42318;--rb:#fdecea;--a:#8a5a00;--ab:#fff3d6;--g:#1d7445;--gb:#e5f3ea;--b:#38598a;--bb:#e8eef8}
@media (prefers-color-scheme:dark){:root{--bg:#11161d;--card:#18202a;--ink:#e6ebf1;--muted:#9aa6b5;--line:#2a3442;
--accent:#7db4dc;--r:#ff9b8f;--rb:#3a1d1b;--a:#f2c46b;--ab:#382a12;--g:#7fd3a0;--gb:#16301f;--b:#a9c1e8;--bb:#1d2a3d}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
main{max-width:1120px;margin:0 auto;padding:28px 16px 64px}
.eyebrow{color:var(--accent);font-weight:600;font-size:12px;letter-spacing:.06em;text-transform:uppercase}
h1{font-size:26px;line-height:1.25;margin:4px 0 6px}h2{font-size:17px;margin:0 0 2px}
.sub,.lead{color:var(--muted);font-size:14px;margin:0 0 14px}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:12px;margin-top:18px}
.kpi,.passo{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px 14px}
.kpi span,.passo span{color:var(--muted);font-size:12.5px}.kpi b{display:block;font-size:19px;white-space:nowrap;font-variant-numeric:tabular-nums}
.passo b{display:block;font-size:14px}.passo small{color:var(--muted)}
section{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:18px 20px;margin-top:16px}
.wrap{overflow-x:auto}table{width:100%;border-collapse:collapse;font-size:13.5px}
th{text-align:left;color:var(--muted);font-weight:600;font-size:11.5px;text-transform:uppercase;letter-spacing:.04em;
border-bottom:1px solid var(--line);padding:7px 8px}td{border-bottom:1px solid var(--line);padding:7px 8px;vertical-align:top}
td.n,th.n{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
tr.det td{color:var(--muted);font-size:12.5px;padding-top:0}tr.lin td{border-bottom:0}
h3{font-size:14px;margin:18px 0 6px;display:flex;gap:8px;align-items:center}
.pill{display:inline-block;border-radius:999px;padding:1px 9px;font-size:12px;font-weight:600;white-space:nowrap}
.r{background:var(--rb);color:var(--r)}.a{background:var(--ab);color:var(--a)}.g{background:var(--gb);color:var(--g)}.b{background:var(--bb);color:var(--b)}
code{font-size:12px;color:var(--muted)}ul{margin:6px 0 0;padding-left:18px}li{margin:2px 0}
footer{color:var(--muted);font-size:12.5px;margin-top:20px}
"""


def montar(ctx) -> dict:
    log = ctx.gw.log_path
    return {"ciclo": ctx.ciclo, "status": ctx.reg.status(ctx.ciclo), "fontes": ctx.reg.fontes(ctx.ciclo) or ctx.fontes(),
            "parametros": ctx.reg.parametros(ctx.ciclo), "ex": ctx.reg.ultima_execucao(ctx.ciclo),
            "aprovacoes": ctx.reg.aprovacoes(ctx.ciclo), "eventos": ctx.reg.eventos(ctx.ciclo), "faixas": ctx.faixas,
            "gateway": [json.loads(l) for l in log.read_text(encoding="utf-8").splitlines()] if log.exists() else []}


def situacao(d, fechado) -> tuple[str, str]:
    if d["faixa"] == "classe própria":
        return ("registrado · ANOMALIA: conferir medição", "a") if d.get("alerta") else ("registrado (esperado)", "b")
    if d["faixa"] == "dentro da faixa":
        return "passa com regra registrada", "g"
    dec = d.get("decisao")
    if dec:
        return f"{dec['decisao']} por {dec['por']} ({dec['papel']}) em {dec['data']}", PILL[dec["decisao"]]
    return ("seguiu sem permissão (sinalizado)" if fechado else "pendente de permissão"), "r"


def pill(texto, cls=None):
    return f'<span class="pill {cls or PILL.get(texto, "b")}">{escape(texto)}</span>'


def tabela(cab, linhas, num_cols=()):
    th = "".join(f'<th class="n">{escape(c)}</th>' if i in num_cols else f"<th>{escape(c)}</th>" for i, c in enumerate(cab))
    corpo = "".join("<tr>" + "".join(f'<td class="n">{c}</td>' if i in num_cols else f"<td>{c}</td>"
                                     for i, c in enumerate(l)) + "</tr>" for l in linhas)
    return f'<div class="wrap"><table><thead><tr>{th}</tr></thead><tbody>{corpo}</tbody></table></div>'


def html(dados: dict) -> str:
    ex, e = dados["ex"], escape
    res, difs = ex["resultado"], ex["diferencas"]
    fechado = dados["status"]["status"] == "fechado"
    ano, mes = dados["ciclo"].split("-")
    tot_f = sum(D(t["total"]) for t in res["totais"].values())
    tot_e = sum(D(v) for v in res["erp_totais"].values())
    acima = [d for d in difs if d["faixa"] in ("acima da faixa", "acima da materialidade")]
    sem_erp = res.get("erp_ausente", False)
    pendentes = sum(p["status"] == "pendente" for p in dados["parametros"])
    permitidas = [d for d in acima if d.get("decisao") and d["decisao"]["decisao"] == "permitida"]
    fontes_ok = sum(f["status"] == "recebida" for f in dados["fontes"])
    modos = sorted({g["modo"] for g in dados["gateway"]}) or ["sem chamadas"]
    aceitos = sum(p["status"] == "aceito" for p in dados["parametros"])
    explicadas = sum(bool(d.get("explicacao")) for d in acima)

    kpis = [("Apurado pela ferramenta", brl(tot_f)), ("Pré-fatura do ERP", "não chegou" if sem_erp else brl(tot_e)),
            ("Diferença líquida (ferramenta − ERP)", "—" if sem_erp else brl(tot_f - tot_e)),
            ("Acima da faixa: permitidas", f"{len(permitidas)} de {len(acima)}"),
            ("Fontes completas", f"{fontes_ok} de {len(dados['fontes'])}")]
    passos = [("1 · Reúne", "agente", f"{fontes_ok} de {len(dados['fontes'])} fontes completas"),
              ("2 · Parametriza", "agente propõe, pessoa aceita",
               f"{len(dados['parametros'])} propostos, {aceitos} aceitos" + (f", {pendentes} PENDENTES" if pendentes else "")),
              ("3 · Calcula", "código determinístico", f"versão {ex['versao']}"),
              ("4 · Explica", "agente", f"{explicadas} de {len(acima)} explicadas ({', '.join(modos)})"),
              ("5 · Aprova", "pessoa", "ciclo fechado" if fechado else "ciclo aberto")]
    out = [f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Ciclo {e(dados['ciclo'])} · Transportadora Exemplo</title><style>{CSS}</style></head><body><main>
<div class="eyebrow">Protótipo · dados 100% sintéticos</div>
<h1>Ciclo de custos de transporte · {MESES[int(mes) - 1]} de {ano}</h1>
<p class="sub">Transportadora Exemplo (fictícia) · {pill(dados['status']['status'])} · versão <code>{e(ex['versao'])}</code> ·
saídas <code>{e(ex['hash_saidas'])}</code> · agente em modo {e(', '.join(modos))} · regras {e(dados['faixas']['versao'])} ({e(dados['faixas']['modo'])})</p>
<div class="grid">{''.join(f'<div class="kpi"><span>{e(k)}</span><b>{e(v)}</b></div>' for k, v in kpis)}</div>
<div class="grid">{''.join(f'<div class="passo"><b>{e(a)}</b><span>{e(b)}</span><br><small>{e(c)}</small></div>' for a, b, c in passos)}</div>"""]

    out.append('<section><h2>1. Fontes do mês</h2><p class="lead">O agente confere o manifesto: o que chegou, o que falta e de quem cobrar.</p>')
    out.append(tabela(["Fonte", "Origem", "Responsável", "Arquivo", "Situação", "Observação"],
                      [[e(f["nome"]), e(f["fonte"]), e(f["responsavel"]), f"<code>{e(f['arquivo'])}</code>", pill(f["status"]), e(f["obs"])]
                       for f in dados["fontes"]]) + "</section>")

    out.append('<section><h2>2. Parâmetros propostos pelo agente</h2><p class="lead">Lidos de e-mails sem estrutura; '
               'nenhum entra no cálculo antes de uma pessoa aceitar, e o ciclo não fecha com proposta sem decisão. '
               'O valor e a data de vigência são conferidos no texto de origem.</p>')
    out.append(tabela(["ID", "Para", "Parâmetro", "Vigente / referência", "Proposto", "Vigência", "Origem", "Situação"],
                      [[f"P{p['id']}", e(NOMES.get(p["carregador"], p["carregador"])), e(PARAM_NOMES.get(p["parametro"], p["parametro"])),
                        brl(p["valor_anterior"]) if p["valor_anterior"] else "—", brl(p["valor"]), e(p["vigencia_inicio"]),
                        e(p["origem"]) + ("" if p["verificado"] else " ⚠ valor ou vigência não encontrados no texto"),
                        pill(p["status"]) + (f"<br><small>{e(p['decidido_por'] or '')}</small>" if p["decidido_por"] else "")]
                       for p in dados["parametros"]], num_cols=(3, 4)) + "</section>")

    componentes = [("capacidade", "Capacidade"), ("movimentacao", "Movimentação"), ("gas_uso_sistema", "Gás de uso"),
                   ("capacidade_extraordinaria", "Cap. extra")]  # coluna toda zerada sai da tabela (fica no XLSX)
    componentes = [(k, n) for k, n in componentes if any(D(t[k]) for t in res["totais"].values())]
    linhas = []
    for car, t in res["totais"].items():
        base = res["volumes"][car]
        erp = D(res["erp_totais"].get(car, 0))
        linhas.append([e(NOMES[car]) + ("" if base["final"] else '<br><span class="pill a">base provisória</span>'),
                       *(brl(t[k]) for k, _ in componentes),
                       brl(t["total"]), "—" if sem_erp else brl(erp), "—" if sem_erp else f"<b>{brl(D(t['total']) - erp)}</b>"])
    out.append('<section><h2>3. Totais por carregador</h2><p class="lead">Calculados pelo código, sobre a medição final '
               f'certificada quando existe. Gás de uso do sistema: {num(D(res["gus"]["consumo_mil_m3"]), 3)} mil m³ a '
               f'{brl(res["gus"]["preco_usado"])}/mil m³ ({e(res["gus"]["origem_preco"])}), rateado pelo volume transportado.'
               + ("" if any(k == "capacidade_extraordinaria" for k, _ in componentes) else
                  " Capacidade extraordinária: nada apurado pela ferramenta neste ciclo (coluna omitida).") + '</p>')
    cab = ["Carregador", *(n for _, n in componentes), "Total", "ERP", "Diferença"]
    out.append(tabela(cab, linhas, num_cols=tuple(range(1, len(cab)))) + "</section>")

    out.append('<section><h2>4. Diferenças por faixa de tolerância</h2><p class="lead">Ferramenta × ERP, decompostas em efeito '
               'de quantidade, de tarifa e de aritmética do ERP (somam exatamente a diferença líquida), e ferramenta × esperado '
               f'(preço acima do índice, fora dessa soma). Linha de materialidade: {brl(dados["faixas"]["materialidade_brl"])}. '
               f'{e(dados["faixas"]["nota_modo"])}</p>')
    for faixa in FAIXAS:
        grupo = [d for d in difs if d["faixa"] == faixa]
        if not grupo:
            continue
        out.append(f"<h3>{pill(faixa)} {len(grupo)} item(ns) · líquido {brl(sum(D(d['valor']) for d in grupo))}</h3>")
        corpo = ""
        for d in grupo:
            txt, cls = situacao(d, fechado)
            corpo += (f'<tr class="lin"><td><b>{d["id"]}</b></td><td>{e(NOMES.get(d["carregador"], d["carregador"]))} · '
                      f'{e(COMPONENTES.get(d["componente"], d["componente"]))}</td><td>{e(d["classe"])}</td>'
                      f'<td class="n">{brl(d["valor"])}</td><td class="n">{d["idade"]}</td><td>{pill(txt, cls)}</td></tr>')
            det = f"<b>Causa provável:</b> {e(d['causa'])} <b>Fonte:</b> {e(d['fonte'])}."
            if d.get("alerta"):
                det += f"<br><b>Anomalia:</b> {e(d['alerta'])}"
            if d.get("explicacao"):
                det += f"<br><b>Explicação do agente ({e(d['explicacao_modo'])}):</b> {e(d['explicacao'])}"
            if d.get("decisao"):
                det += f"<br><b>Motivo registrado:</b> {e(d['decisao']['motivo'])}"
            elif faixa in ("dentro da faixa", "classe própria"):
                det += f"<br><b>Regra:</b> {e(d['regra'])}"
            corpo += f'<tr class="det"><td></td><td colspan="5">{det}</td></tr>'
        out.append('<div class="wrap"><table><thead><tr><th>ID</th><th>Carregador · componente</th><th>Classe</th>'
                   f'<th class="n">Valor</th><th class="n">Idade (ciclos)</th><th>Situação</th></tr></thead><tbody>{corpo}</tbody></table></div>')
    out.append("</section>")

    por_chave = {d["chave"]: d["id"] for d in difs}
    out.append('<section><h2>5. Permissões e devoluções</h2><p class="lead">Acima da faixa, cada item é permitido por escrito '
               'por alguém com nome e papel, ou devolvido. A permissão vale para aquele valor: se o valor mudar, vale de novo.</p>')
    out.append(tabela(["Data", "Diferença", "Valor", "Decisão", "Quem", "Motivo"],
                      [[e(a["data"]), por_chave.get(a["chave"], e(a["chave"])), brl(a["valor"]), pill(a["decisao"]),
                        f"{e(a['por'])}<br><small>{e(a['papel'])}</small>", e(a["motivo"])] for a in dados["aprovacoes"]],
                      num_cols=(2,)) + "</section>")

    partes = ex["versao_partes"]
    out.append('<section><h2>6. Trilha de auditoria</h2><p class="lead">Mesmas entradas, regras e parâmetros aceitos '
               'produzem a mesma versão e as mesmas saídas.</p><ul>' +
               f"<li>Versão <code>{e(ex['versao'])}</code> = cálculo <code>{e(partes['calculo'])}</code> (código sha256 "
               f"<code>{e(partes['codigo_sha256'][:16])}…</code>) + regras <code>{e(partes['regras'])}</code> + "
               f"{len(partes['parametros'])} parâmetro(s) aceito(s) + {len(partes['referencias'])} arquivo(s) de referência + "
               f"{len(partes['entradas'])} arquivo(s) de entrada do mês:</li>" +
               "".join(f"<li><code>{e(k)}</code> sha256 <code>{v[:16]}…</code></li>"
                       for k, v in sorted({**partes["referencias"], **partes["entradas"]}.items())) +
               "</ul><h3>Chamadas ao gateway de IA</h3>")
    out.append(tabela(["Hora", "Finalidade", "Modo", "Bytes", "Dados pessoais cortados", "sha256 do conteúdo"],
                      [[e(g["ts"][11:]), e(g["finalidade"]), e(g["modo"]), str(g["bytes"]),
                        e(", ".join(f"{k}: {v}" for k, v in g["cortes"].items()) or "nenhum"), f"<code>{g['sha256'][:12]}…</code>"]
                       for g in dados["gateway"]], num_cols=(3,)))
    out.append("<h3>Eventos do ciclo</h3><ul>" + "".join(
        f"<li><code>{e(v['ts'][5:16].replace('T', ' '))}</code> <b>{e(v['tipo'])}</b>: {e(v['detalhe'])}</li>" for v in dados["eventos"]) + "</ul></section>")
    out.append(f"<footer>Gerado em {agora():%d/%m/%Y %H:%M} pelo protótipo <code>prototipo-ciclo-agentico</code>. "
               "Empresas, pessoas, volumes, tarifas e preços são inventados.</footer></main></body></html>")
    return "\n".join(out)


def xlsx(dados: dict, caminho: Path) -> None:
    ex = dados["ex"]
    res = ex["resultado"]
    wb = Workbook()
    abas = {
        "Totais": (["Carregador", "Base de medição", "Capacidade", "Movimentação", "Gás de uso do sistema", "Capacidade extraordinária",
                    "Total ferramenta", "Pré-fatura ERP", "Diferença"],
                   [[NOMES[c], "final" if res["volumes"][c]["final"] else "provisória",
                     *(float(t[k]) for k in ("capacidade", "movimentacao", "gas_uso_sistema", "capacidade_extraordinaria", "total")),
                     *(["não chegou", None] if res.get("erp_ausente") else
                       [float(res["erp_totais"].get(c, 0)), float(D(t["total"]) - D(res["erp_totais"].get(c, 0)))])]
                    for c, t in res["totais"].items()]),
        "Linhas calculadas": (["Carregador", "Componente", "Quantidade", "Unidade", "Tarifa", "Valor", "Origem da tarifa", "Base"],
                              [[NOMES[l["carregador"]], COMPONENTES[l["componente"]], float(l["quantidade"]), l["unidade"], float(l["tarifa"]),
                                float(l["valor"]), l["origem_tarifa"], l["base_medicao"]] for l in res["linhas"]]),
        "Diferenças": (["ID", "Carregador", "Componente", "Classe", "Faixa", "Valor", "Idade (ciclos)", "Causa provável", "Fonte",
                        "Anomalia", "Explicação", "Situação"],
                       [[d["id"], NOMES.get(d["carregador"], d["carregador"]), COMPONENTES.get(d["componente"], d["componente"]), d["classe"], d["faixa"],
                         float(d["valor"]), d["idade"], d["causa"], d["fonte"], d.get("alerta") or "", d.get("explicacao") or "",
                         situacao(d, dados["status"]["status"] == "fechado")[0]] for d in ex["diferencas"]]),
        "Permissões": (["Data", "Chave", "Valor", "Decisão", "Quem", "Papel", "Motivo"],
                       [[a["data"], a["chave"], float(a["valor"]), a["decisao"], a["por"], a["papel"], a["motivo"]] for a in dados["aprovacoes"]]),
        "Fontes": (["Fonte", "Responsável", "Arquivo", "Situação", "Observação"],
                   [[f["nome"], f["responsavel"], f["arquivo"], f["status"], f["obs"]] for f in dados["fontes"]]),
        "Trilha": (["Item", "Valor"], [["Ciclo", dados["ciclo"]], ["Situação", dados["status"]["status"]], ["Versão", ex["versao"]],
                                       ["Hash das saídas", ex["hash_saidas"]], ["Regras", ex["versao_partes"]["regras"]]] +
                   [["sha256 do código de cálculo", ex["versao_partes"]["codigo_sha256"]]] +
                   [[f"sha256 {k}", v] for k, v in sorted({**ex["versao_partes"]["referencias"], **ex["versao_partes"]["entradas"]}.items())] +
                   [[f"{v['ts']} {v['tipo']}", v["detalhe"]] for v in dados["eventos"]]),
    }
    wb.remove(wb.active)
    for nome, (cab, linhas) in abas.items():
        ws = wb.create_sheet(nome)
        ws.append(cab)
        for l in linhas:
            ws.append(l)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1D5F8A")
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, start=1):
            largura = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, largura + 2), 60)
            for c in col[1:]:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00" if "Quantidade" not in str(col[0].value) else "#,##0.000"
    wb.properties.creator = "prototipo-ciclo-agentico"
    wb.properties.created = wb.properties.modified = agora()
    wb.save(caminho)
    if relogio_fixo():
        _zip_estavel(caminho)


def _zip_estavel(caminho: Path) -> None:
    """Regrava o XLSX com data fixa nas entradas do zip: o snapshot do demo fica idêntico entre execuções."""
    with zipfile.ZipFile(caminho) as z:
        itens = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    criado = re.compile(rb"(<dcterms:created[^>]*>)([^<]+)")  # o openpyxl grava 'modified' com a hora real ao salvar
    itens = [(n, re.sub(rb"(<dcterms:modified[^>]*>)[^<]+", lambda m: m.group(1) + criado.search(d).group(2), d)
              if n == "docProps/core.xml" else d) for n, d in itens]
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, dado in itens:
            z.writestr(zipfile.ZipInfo(nome, date_time=(2026, 10, 7, 9, 0, 0)), dado, zipfile.ZIP_DEFLATED)


def gerar(ctx) -> tuple[Path, Path]:
    dados = montar(ctx)
    if not dados["ex"]:
        raise ValueError("Nenhum cálculo registrado: rode 'calcular' primeiro.")
    p_html, p_xlsx = ctx.saida / f"relatorio_{ctx.ciclo}.html", ctx.saida / f"ciclo_{ctx.ciclo}.xlsx"
    p_html.write_text(html(dados), encoding="utf-8")
    xlsx(dados, p_xlsx)
    return p_html, p_xlsx
