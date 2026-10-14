"""Utilitários: dinheiro em Decimal, formatação pt-BR, hashes e leitura de CSV."""
import csv
import hashlib
import json
from datetime import datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

_RELOGIO = []  # vazio = relógio real; [datetime] = relógio fixo do demo (avança 1 s por leitura)


def fixar_relogio(inicio: datetime | None) -> None:
    """O demo fixa o relógio para que o snapshot em exemplo/ seja idêntico a cada execução."""
    _RELOGIO[:] = [inicio] if inicio else []


def relogio_fixo() -> bool:
    return bool(_RELOGIO)


def agora() -> datetime:
    if _RELOGIO:
        _RELOGIO[0] += timedelta(seconds=1)
        return _RELOGIO[0]
    return datetime.now().replace(microsecond=0)


def agora_iso() -> str:
    return agora().isoformat(timespec="seconds")


def D(x) -> Decimal:
    return x if isinstance(x, Decimal) else Decimal(str(x))


def q2(x) -> Decimal:
    """Dinheiro: 2 casas, arredondamento comercial."""
    return D(x).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def q3(x) -> Decimal:
    """Volume (mil m³): 3 casas."""
    return D(x).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)


def num(x, casas=2) -> str:
    """1234567.891 -> '1.234.567,89' (pt-BR)."""
    s = f"{D(x):,.{casas}f}"
    return s.replace(",", "_").replace(".", ",").replace("_", ".")


def brl(x) -> str:
    v = D(x)
    return ("-R$ " if v < 0 else "R$ ") + num(abs(v))


def mes_indice(ciclo: str) -> int:
    ano, mes = ciclo.split("-")
    return int(ano) * 12 + int(mes)


def canon(obj) -> str:
    return json.dumps(obj, sort_keys=True, ensure_ascii=False, default=str, separators=(",", ":"))


def sha256_texto(texto: str) -> str:
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def sha256_arquivo(caminho: Path) -> str:
    caminho = Path(caminho)
    if caminho.suffix == ".xlsx":  # conteúdo das células, não os metadados que mudam a cada salvamento
        from openpyxl import load_workbook
        wb = load_workbook(caminho, read_only=True)
        conteudo = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}
        wb.close()
        return sha256_texto(canon(conteudo))
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def sha256_codigo(caminhos) -> str:
    """Hash do código-fonte com fim de linha normalizado (o mesmo em Windows e Linux)."""
    return sha256_texto("".join(Path(c).read_text(encoding="utf-8").replace("\r\n", "\n") for c in caminhos))


def ler_csv(caminho: Path) -> list[dict]:
    with open(caminho, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def escrever_csv(caminho: Path, cabecalho: list[str], linhas: list[list]) -> None:
    with open(caminho, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(cabecalho)
        w.writerows(linhas)
