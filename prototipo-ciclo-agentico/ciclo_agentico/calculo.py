"""Cálculo determinístico: nenhum valor em dinheiro sai do modelo, só deste código.

Encargo de capacidade  = capacidade contratada vigente × tarifa de capacidade
Encargo de movimentação = volume do mês (final certificada; se faltar, provisória sinalizada) × tarifa
Gás de uso do sistema   = consumo do mês × preço, rateado pelo volume transportado de cada carregador
Capacidade extraordinária = quantidade do portal × tarifa (só se a fonte chegou)
"""
from pathlib import Path

from openpyxl import load_workbook

from .util import D, canon, ler_csv, q2, q3, sha256_codigo, sha256_texto

VERSAO_CALCULO = "calc-v1"


CODIGO = [Path(__file__).with_name(n) for n in ("calculo.py", "comparacao.py", "util.py")]


def ler_contratos(caminho: Path) -> list[dict]:
    wb = load_workbook(caminho, read_only=True)
    try:
        linhas = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    return [dict(zip(linhas[0], l)) for l in linhas[1:] if l[0]]


def ler_indice(caminho: Path) -> dict:
    return {r["mes"]: D(r["preco_referencia_brl_mil_m3"]) for r in ler_csv(caminho)}


def carregar_entradas(dir_ciclo: Path) -> dict:
    """Só é chamada depois de conferir que as fontes que bloqueiam o cálculo chegaram (operacoes.calcular)."""
    opcional = lambda nome: ler_csv(dir_ciclo / nome) if (dir_ciclo / nome).exists() else None
    return {
        "prov": ler_csv(dir_ciclo / "medicao_provisoria.csv"),
        "final": opcional("medicao_final.csv") or [],
        "gus": ler_csv(dir_ciclo / "consumo_gas_sistema.csv"),
        "indice": ler_indice(dir_ciclo / "indice_preco_gas.csv"),
        "contratos": ler_contratos(dir_ciclo / "contratos.xlsx"),
        "extra": opcional("capacidade_extraordinaria.csv"),
        "erp": opcional("pre_fatura_erp.csv") or [],
    }


def _somar(linhas) -> dict:
    tot = {}
    for l in linhas:
        tot[l["carregador"]] = tot.get(l["carregador"], D(0)) + D(l["volume_mil_m3"])
    return {c: q3(v) for c, v in tot.items()}


def _param(params, carregador, nome, ciclo):
    vale = [p for p in params if p["carregador"] == carregador and p["parametro"] == nome and p["vigencia_inicio"][:7] <= ciclo]
    return max(vale, key=lambda p: p["vigencia_inicio"]) if vale else None


def calcular(ent: dict, params: list[dict], ciclo: str) -> dict:
    inicio = f"{ciclo}-01"
    vol_prov, vol_final = _somar(ent["prov"]), _somar(ent["final"])
    vol = {c: vol_final.get(c, v) for c, v in vol_prov.items()}
    tot, tot_prov = sum(vol.values()), sum(vol_prov.values())
    gus_total = q3(sum(D(r["consumo_mil_m3"]) for r in ent["gus"]))
    preco_ref = ent["indice"][ciclo]
    p_gus = _param(params, "SISTEMA", "preco_gas_uso_sistema", ciclo)
    preco_gus = D(p_gus["valor"]) if p_gus else preco_ref
    origem_gus = f"parâmetro aceito, {p_gus['origem']}" if p_gus else "índice público do mês (sem preço aceito do fornecedor)"

    linhas = []

    def linha(car, comp, qtd, unid, tarifa, origem, q_prov=None):
        linhas.append({"carregador": car, "componente": comp, "quantidade": str(qtd), "unidade": unid,
                       "tarifa": str(q2(tarifa)), "valor": str(q2(D(qtd) * q2(tarifa))), "origem_tarifa": origem,
                       "quantidade_provisoria": str(qtd if q_prov is None else q_prov),
                       "base_medicao": "final" if car in vol_final else "provisória"})

    for ct in ent["contratos"]:
        car = ct["carregador"]
        if not (str(ct["vigencia_inicio"]) <= inicio <= str(ct["vigencia_fim"])) or car not in vol:
            continue
        for comp, nome, unid, qtd, q_prov in (
                ("capacidade", "tarifa_capacidade", "mil m³/dia", D(ct["capacidade_mil_m3_dia"]), None),
                ("movimentacao", "tarifa_movimentacao", "mil m³", vol[car], vol_prov[car])):
            p = _param(params, car, nome, ciclo)
            tarifa = D(p["valor"]) if p else D(str(ct[f"{nome}_brl"]))
            linha(car, comp, qtd, unid, tarifa, f"parâmetro aceito, {p['origem']}" if p else f"contrato {ct['contrato']}", q_prov)
        linha(car, "gas_uso_sistema", q3(gus_total * vol[car] / tot), "mil m³", preco_gus, origem_gus,
              q3(gus_total * vol_prov[car] / tot_prov))
        for r in ent["extra"] or []:
            if r["carregador"] == car:
                linha(car, "capacidade_extraordinaria", D(r["quantidade_mil_m3_dia"]), "mil m³/dia", D(r["tarifa_brl"]), "portal de capacidade")

    totais = {}
    for l in linhas:
        t = totais.setdefault(l["carregador"], {"capacidade": D(0), "movimentacao": D(0), "gas_uso_sistema": D(0),
                                                 "capacidade_extraordinaria": D(0), "total": D(0)})
        t[l["componente"]] += D(l["valor"])
        t["total"] += D(l["valor"])
    meses = sorted(m for m in ent["indice"] if m < ciclo)
    return {
        "ciclo": ciclo, "versao_calculo": VERSAO_CALCULO, "linhas": linhas,
        "totais": {c: {k: str(v) for k, v in t.items()} for c, t in totais.items()},
        "volumes": {c: {"provisoria": str(vol_prov[c]), "final": str(vol_final.get(c, "")), "usado": str(vol[c])} for c in vol},
        "gus": {"consumo_mil_m3": str(gus_total), "preco_usado": str(q2(preco_gus)), "origem_preco": origem_gus,
                "preco_esperado": str(q2(preco_ref)), "preco_mes_anterior": str(ent["indice"][meses[-1]]) if meses else ""},
    }


def versao(hashes_entradas: dict, referencias: dict, faixas: dict, params: list[dict]) -> dict:
    """Mesmas entradas do mês + mesmas referências (histórico, manifesto, contatos) + mesmas regras +
    mesmo código de cálculo + mesmos parâmetros aceitos => mesma versão."""
    partes = {"entradas": hashes_entradas, "referencias": referencias, "regras": faixas["versao"],
              "regras_sha256": sha256_texto(canon(faixas)), "calculo": VERSAO_CALCULO, "codigo_sha256": sha256_codigo(CODIGO),
              "parametros": sorted(canon({k: p[k] for k in ("carregador", "parametro", "valor", "vigencia_inicio")}) for p in params)}
    return {"hash": sha256_texto(canon(partes))[:16], "partes": partes}
